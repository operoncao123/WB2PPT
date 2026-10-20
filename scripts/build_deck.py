#!/usr/bin/env python3
"""Render spec.yaml -> native editable PPTX + HTML preview with identical geometry.

Usage:
  python3 build_deck.py --spec spec.yaml --ppt wb_demo.pptx --html wb_demo.html

See SKILL.md for the full spec schema. Geometry summary:
- blots stacked at layout.blot_x, width layout.blot_w, starting layout.y0, gap layout.gap;
  heights come from each blot's native image aspect (h optional override per blot).
- MW label right-aligned in a box ending 0.18in left of the panel edge; tick line just
  left of the panel at y = blot_y + tick_frac * blot_h.
- sample names (if present): rotated -45deg above the FIRST blot, one per lane, starting
  at each lane center (lane_fracs), fully outside the image.
- antibody right labels at blot right edge + 0.15in, vertically centered on the panel.
- info block at layout-info position (bottom right), plain lines.
"""

import argparse
import base64
import html as html_mod

import yaml
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

IN = 96  # px per inch for the HTML mirror


def build_pptx(spec, out):
    prs = Presentation()
    prs.slide_width = Inches(spec["slide"]["w"])
    prs.slide_height = Inches(spec["slide"]["h"])
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    font = spec.get("font", "Arial")
    unit = spec.get("unit", "")

    def txt(x, y, w, h, lines, size=12, bold=False, align=PP_ALIGN.LEFT, rot=0):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        box.rotation = rot
        tf = box.text_frame
        tf.word_wrap = False if rot else True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        if isinstance(lines, str):
            lines = [lines]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = line
            r.font.name = font
            r.font.size = Pt(size)
            r.font.bold = bold
        return box

    y = 0.10
    if spec.get("title"):
        txt(0.44, y, spec["slide"]["w"] - 0.9, 0.32, spec["title"], size=14, bold=True)
        y += 0.42
    if spec.get("aim"):
        txt(0.44, y, spec["slide"]["w"] - 0.9, 0.30, "Aim: " + spec["aim"], size=13)
        y += 0.34
    if spec.get("conditions"):
        txt(0.44, y, 9.8, 0.30 + 0.24 * len(spec["conditions"]),
            ["Experimental Conditions:"] + [f"\u2022  {c}" for c in spec["conditions"]], size=12)
        y += 0.40 + 0.24 * len(spec["conditions"])
    if spec.get("conclusion"):
        txt(0.44, y + 0.05, 9.8, 0.30 + 0.24 * len(spec["conclusion"]),
            ["Conclusion:"] + [f"\u2022  {c}" for c in spec["conclusion"]], size=12)

    lay = spec["layout"]
    bx, bw, gap = lay["blot_x"], lay["blot_w"], lay.get("gap", 0.06)
    border = lay.get("border_pt", 0.5)
    by = lay["y0"]

    for k, blot in enumerate(spec["blots"]):
        with Image.open(blot["image"]) as im:
            aspect = im.height / im.width
        h = blot.get("h", bw * aspect)
        pic = slide.shapes.add_picture(blot["image"], Inches(bx), Inches(by),
                                       width=Inches(bw), height=Inches(h))
        if border:
            pic.line.color.rgb = RGBColor(0, 0, 0)
            pic.line.width = Pt(border)
        # MW label + tick
        if blot.get("mw"):
            txt(lay.get("mw_label_x", bx - 1.0), by + blot["tick_frac"] * h - 0.14,
                0.82, 0.28, f"{blot['mw']} {unit}".strip(), size=spec.get("label_size", 12),
                align=PP_ALIGN.RIGHT)
            ln = slide.shapes.add_shape(1, Inches(bx - 0.14), Inches(by + blot["tick_frac"] * h),
                                        Inches(0.10), Inches(0.015))
            ln.fill.solid()
            ln.fill.fore_color.rgb = RGBColor(0, 0, 0)
            ln.line.fill.background()
            ln.shadow.inherit = False
            # python-pptx adds <p:style> whose effectRef points at the theme shadow;
            # WPS/LibreOffice render it even with an empty <a:effectLst/> — drop the style.
            style_el = ln._element.find(
                "{http://schemas.openxmlformats.org/presentationml/2006/main}style")
            if style_el is not None:
                ln._element.remove(style_el)
        # antibody right label
        if blot.get("right_label"):
            txt(bx + bw + 0.15, by + h / 2 - 0.15, 1.8, 0.30, blot["right_label"],
                size=spec.get("label_size", 12))
        by += h + gap

    # sample names above the FIRST blot only
    if spec.get("samples") and spec.get("lane_fracs"):
        top_y = lay["y0"]
        for frac, nm in zip(spec["lane_fracs"], spec["samples"]):
            lx = bx + frac * bw
            txt(lx - 0.135, top_y - 0.50, 0.9, 0.25, nm, size=spec.get("label_size", 12), rot=315)

    info = spec.get("info")
    if info:
        txt(info["x"], info["y"], info["w"], 0.20 * len(info["lines"]) + 0.1,
            info["lines"], size=info.get("size", 11))

    prs.save(out)


def build_html(spec, out):
    def pos(x, y, w, h):
        return f"left:{x * IN:.0f}px;top:{y * IN:.0f}px;width:{w * IN:.0f}px;height:{h * IN:.0f}px;"

    def b64(p):
        return base64.b64encode(open(p, "rb").read()).decode()

    font = spec.get("font", "Arial")
    unit = spec.get("unit", "")
    lay = spec["layout"]
    bx, bw, gap = lay["blot_x"], lay["blot_w"], lay.get("gap", 0.06)
    border = int(round(lay.get("border_pt", 0.5) / 72 * IN)) or 1

    el = []

    def add(txt_html, style, cls="txt"):
        el.append(f'  <div class="{cls}" style="{style}{style_extra}">{txt_html}</div>')

    style_extra = ""
    y = 0.10
    if spec.get("title"):
        el.append(f'  <div class="txt" style="{pos(0.44, y, spec["slide"]["w"] - 0.9, 0.32)} font-size:19px; font-weight:bold;">{html_mod.escape(spec["title"])}</div>')
        y += 0.42
    if spec.get("aim"):
        el.append(f'  <div class="txt" style="{pos(0.44, y, spec["slide"]["w"] - 0.9, 0.30)} font-size:17px;">Aim: {html_mod.escape(spec["aim"])}</div>')
        y += 0.34
    if spec.get("conditions"):
        lines = "<br>".join(html_mod.escape(c) for c in spec["conditions"])
        el.append(f'  <div class="txt" style="{pos(0.44, y, 9.8, 0.30 + 0.24 * len(spec["conditions"]))} font-size:16px; line-height:1.35;">Experimental Conditions:<br>&bull;&nbsp; {lines}</div>')
        y += 0.40 + 0.24 * len(spec["conditions"])
    if spec.get("conclusion"):
        lines = "<br>&bull;&nbsp; ".join(html_mod.escape(c) for c in spec["conclusion"])
        el.append(f'  <div class="txt" style="{pos(0.44, y + 0.05, 9.8, 0.30 + 0.24 * len(spec["conclusion"]))} font-size:16px; line-height:1.35;">Conclusion:<br>&bull;&nbsp; {lines}</div>')

    by = lay["y0"]
    for k, blot in enumerate(spec["blots"]):
        with Image.open(blot["image"]) as im:
            aspect = im.height / im.width
        h = blot.get("h", bw * aspect)
        data = b64(blot["image"])
        el.append(f'  <img class="blot" src="data:image/png;base64,{data}" style="{pos(bx, by, bw, h)}">')
        if blot.get("mw"):
            ty = by + blot["tick_frac"] * h
            el.append(f'  <div class="txt mw" style="{pos(lay.get("mw_label_x", bx - 1.0), ty - 0.14, 0.82, 0.28)}">{html_mod.escape(blot["mw"] + " " + unit).strip()}</div>')
            el.append(f'  <div class="tick" style="left:{(bx - 0.14) * IN:.0f}px;top:{ty * IN:.0f}px;width:{0.10 * IN:.0f}px;"></div>')
        if blot.get("right_label"):
            el.append(f'  <div class="txt abname" style="{pos(bx + bw + 0.15, by + h / 2 - 0.15, 1.8, 0.30)}">{html_mod.escape(blot["right_label"])}</div>')
        by += h + gap

    if spec.get("samples") and spec.get("lane_fracs"):
        for frac, nm in zip(spec["lane_fracs"], spec["samples"]):
            lx = bx + frac * bw
            st = pos(lx + 0.051, lay["y0"] - 0.219, 0.9, 0.25) + "transform-origin:left bottom;transform:rotate(-45deg);"
            el.append(f'  <div class="txt sname" style="{st}">{html_mod.escape(nm)}</div>')

    info = spec.get("info")
    if info:
        body = "<br>".join(html_mod.escape(l).replace("  ", "&nbsp;&nbsp;") for l in info["lines"])
        el.append(f'  <div class="txt" style="{pos(info["x"], info["y"], info["w"], 0.20 * len(info["lines"]) + 0.1)} font-size:{info.get("size", 11) * 1.35:.0f}px; line-height:1.3;">{body}</div>')

    doc = f"""<!DOCTYPE html>
<html lang="zh">
<head>
<meta charset="UTF-8">
<title>{html_mod.escape(spec.get("title", "Western Blot"))}</title>
<style>
  body {{ margin:0; background:#888; font-family: {font}, "PingFang SC", "Microsoft YaHei", sans-serif; }}
  #slide {{ position:relative; width:{spec["slide"]["w"] * IN:.0f}px; height:{spec["slide"]["h"] * IN:.0f}px; background:#fff; margin:20px auto; box-shadow:0 4px 18px rgba(0,0,0,.35); overflow:hidden; }}
  .txt {{ position:absolute; color:#000; }}
  .blot {{ position:absolute; border:{border}px solid #000; box-sizing:border-box; }}
  .mw {{ text-align:right; font-size:16px; white-space:nowrap; }}
  .abname {{ font-size:16px; }}
  .sname {{ font-size:16px; white-space:nowrap; }}
  .tick {{ position:absolute; background:#000; height:2px; }}
</style>
</head>
<body>
<div id="slide">
{chr(10).join(el)}
</div>
</body>
</html>
"""
    open(out, "w").write(doc)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--spec", required=True)
    ap.add_argument("--ppt", required=True)
    ap.add_argument("--html", required=True)
    args = ap.parse_args()

    spec = yaml.safe_load(open(args.spec))
    build_pptx(spec, args.ppt)
    build_html(spec, args.html)
    print(f"saved {args.ppt} and {args.html}")


if __name__ == "__main__":
    main()
